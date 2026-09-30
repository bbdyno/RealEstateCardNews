"""정비사업(재개발·재건축) 공개 파일 가져오기. 지자체마다 열 이름이 달라 흔한 이름을 알아서 찾는다.

  .venv/bin/python -m collector.redevelop 파일.csv --sido 서울 --source "서울시 정비사업 데이터" --date 2021-12-27 --dry-run
  .venv/bin/python -m collector.redevelop 파일.csv --sido 대구 --source "대구광역시 정비사업 추진현황" --date 2025-01-31

- CSV(UTF-8 / CP949)와 XLSX 를 읽는다(XLSX 는 openpyxl 필요).
- --dry-run 은 찾은 열과 앞 5줄만 보여 주고 넣지 않는다. 열을 못 찾으면 --map zone=구역명,stage=진행단계 식으로 알려 준다.
- 같은 출처·같은 구역명이면 덮어쓴다(갱신된 파일을 다시 넣으면 단계가 바뀐다).

공개 파일 후보(공공데이터포털): 서울특별시_서울시 정비사업 데이터(15097425), 서울주택도시공사_공공재개발 및 공공재건축
사업 현황(15124798), 대구광역시_정비사업 추진현황(15141940), 경기데이터드림 '일반 정비 사업 추진 현황'.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import re
from pathlib import Path

from . import db

COLS = {
    "zone": ["정비구역명", "구역명", "사업장명", "구역이름", "사업명", "구역"],
    "sgg_name": ["자치구", "자치구명", "시군구명", "시군구", "시군명", "시군", "구"],
    "address": ["구역위치", "위치", "소재지", "사업장위치", "주소", "대표지번", "대지위치"],
    "biz": ["사업구분", "정비유형", "사업유형", "사업종류", "구분", "사업방식"],
    "stage": ["추진단계", "진행단계", "사업단계", "현추진상황", "추진현황", "추진실정", "단계", "진행상황"],
    "area": ["구역면적(㎡)", "구역면적", "정비구역면적", "면적(㎡)", "면적"],
    "households": ["계획세대수", "공급예정세대수", "공급 예정 세대수", "건립세대수", "공급세대수", "세대수", "총세대수"],
}
# 단계 글 → 번호(0 구역지정 … 7 준공). 앞에 있는 규칙이 먼저 걸린다(준공·착공이 '인가' 같은 말보다 우선).
STAGE_RULES = [
    (7, ("준공", "이전고시", "청산", "입주", "해산")),
    (6, ("착공", "공사중", "공사 중")),
    (5, ("이주", "철거")),
    (4, ("관리처분",)),
    (3, ("사업시행",)),
    (2, ("조합설립", "조합 설립", "조합인가")),
    (1, ("추진위",)),
    (0, ("정비구역", "구역지정", "정비계획", "예정구역", "후보지", "선정", "기본계획", "지정")),
]
STAGE_NAMES = ["정비구역 지정", "추진위원회 승인", "조합설립 인가", "사업시행 인가", "관리처분 인가", "이주·철거", "착공", "준공"]


def stage_no(text: str | None) -> int | None:
    t = (text or "").replace(" ", "")
    if not t:
        return None
    for no, keys in STAGE_RULES:
        if any(k.replace(" ", "") in t for k in keys):
            return no
    return None


def biz_label(text: str | None) -> str:
    t = text or ""
    for k in ("공공재개발", "공공재건축", "모아타운", "모아주택", "가로주택", "소규모", "재건축", "재개발", "도시환경", "주거환경"):
        if k in t:
            return {"도시환경": "도시정비형 재개발", "주거환경": "주거환경개선"}.get(k, k)
    return t.strip() or "정비사업"


def read_rows(path: Path) -> list[dict]:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook          # 선택 의존성
        ws = load_workbook(path, read_only=True, data_only=True).active
        it = ws.iter_rows(values_only=True)
        head = [str(h or "").strip() for h in next(it)]
        return [{h: ("" if v is None else str(v)).strip() for h, v in zip(head, row)} for row in it if any(row)]
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp949"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise SystemExit(f"{path.name}: 인코딩을 알 수 없습니다(UTF-8·CP949 아님).")
    return [{(k or "").strip(): (v or "").strip() for k, v in r.items()} for r in csv.DictReader(io.StringIO(text))]


def detect(headers: list[str], overrides: dict[str, str]) -> dict[str, str | None]:
    """우리 필드 → 파일의 열 이름. 정확히 같은 이름을 먼저, 없으면 포함 관계로 찾는다."""
    out: dict[str, str | None] = {}
    norm = {h: re.sub(r"\s+", "", h) for h in headers}
    for field, cands in COLS.items():
        if field in overrides:
            out[field] = overrides[field]
            continue
        hit = next((h for c in cands for h in headers if norm[h] == c.replace(" ", "")), None)
        if hit is None:
            hit = next((h for c in cands for h in headers if len(c) > 1 and c.replace(" ", "") in norm[h]), None)
        out[field] = hit
    return out


DONG = re.compile(r"([가-힣0-9]+(?:동|가|리))(?=[\s\d,·(]|$)")


def find_umd(address: str, sgg_name: str) -> str:
    """'마포구 아현동 699 일대' → '아현동'. 시군구 이름 뒤에서 처음 나오는 동·가·리."""
    rest = address.split(sgg_name, 1)[-1] if sgg_name and sgg_name in address else address
    m = DONG.search(rest)
    return m.group(1) if m else ""


def to_num(v: str | None) -> float | None:
    v = re.sub(r"[^\d.]", "", v or "")
    try:
        return float(v) if v else None
    except ValueError:
        return None


def region_code(sido_short: str, sgg_name: str) -> str | None:
    if not sgg_name:
        return None
    rs = [r for r in db.regions() if r["sido_short"] == sido_short]
    exact = [r for r in rs if r["name"] == sgg_name or r["name"].replace(" ", "") == sgg_name.replace(" ", "")]
    if exact:
        return exact[0]["code"]
    part = [r for r in rs if sgg_name in r["name"] or r["name"].endswith(sgg_name)]
    return part[0]["code"] if len(part) == 1 else None


def convert(rows: list[dict], cols: dict, sido_short: str, source: str, date: str | None) -> list[dict]:
    sido = next((r["sido"] for r in db.regions() if r["sido_short"] == sido_short), sido_short)
    out = []
    for r in rows:
        get = lambda f: r.get(cols[f] or "", "") if cols.get(f) else ""
        zone = get("zone")
        if not zone:
            continue
        sgg_name = get("sgg_name")
        address = get("address")
        if not sgg_name and address:          # 열이 없으면 주소 앞머리에서 시군구를 찾는다
            m = re.search(r"([가-힣]+(?:구|시|군))", address)
            sgg_name = m.group(1) if m else ""
        stage_txt = get("stage")
        no = stage_no(stage_txt)
        hh = to_num(get("households"))
        out.append({
            "id": hashlib.sha1(f"{source}|{sgg_name}|{zone}".encode()).hexdigest()[:12],
            "sido": sido, "sgg": region_code(sido_short, sgg_name) or "", "sgg_name": sgg_name,
            "umd": find_umd(address, sgg_name) if address else "", "zone": zone, "biz": biz_label(get("biz") or zone),
            "stage": stage_txt or (STAGE_NAMES[no] if no is not None else ""), "stage_no": no,
            "area": to_num(get("area")), "households": int(hh) if hh else None, "address": address,
            "source": source, "source_date": date, "demo": 0,
        })
    return out


FIELDS = ("id", "sido", "sgg", "sgg_name", "umd", "zone", "biz", "stage", "stage_no", "area", "households", "address",
          "source", "source_date", "demo")


def save(zones: list[dict]) -> int:
    c = db.connect()
    c.executemany(f"INSERT OR REPLACE INTO redevelop({','.join(FIELDS)}) VALUES({','.join('?' * len(FIELDS))})",
                  [tuple(z[k] for k in FIELDS) for z in zones])
    c.commit()
    return len(zones)


def main() -> None:
    p = argparse.ArgumentParser()
    p.add_argument("file")
    p.add_argument("--sido", required=True, help="시도 약칭(서울·대구·경기 …)")
    p.add_argument("--source", required=True, help="출처 이름(페이지에 그대로 적힌다)")
    p.add_argument("--date", help="자료 기준일 YYYY-MM-DD")
    p.add_argument("--map", default="", help="열 직접 지정: zone=구역명,stage=진행단계")
    p.add_argument("--dry-run", action="store_true")
    a = p.parse_args()
    overrides = dict(kv.split("=", 1) for kv in a.map.split(",") if "=" in kv)
    path = Path(a.file)
    date = a.date or (lambda m: f"{m[1]}-{m[2]}-{m[3]}" if m else None)(re.search(r"(20\d\d)(\d\d)(\d\d)", path.name))
    rows = read_rows(path)
    cols = detect(list(rows[0]) if rows else [], overrides)
    print("찾은 열:", {k: v for k, v in cols.items()})
    missing = [k for k in ("zone", "stage") if not cols.get(k)]
    if missing:
        print(f"필수 열을 못 찾았습니다: {missing} — --map 으로 알려 주세요. 파일의 열: {list(rows[0]) if rows else []}")
    zones = convert(rows, cols, a.sido, a.source, date)
    unmatched = sorted({z["sgg_name"] for z in zones if not z["sgg"]})
    print(f"{len(zones)}개 구역 · 단계 번호를 못 붙인 구역 {sum(1 for z in zones if z['stage_no'] is None)} · 시군구를 못 찾은 이름 {unmatched[:10]}")
    for z in zones[:5]:
        print("  ", {k: z[k] for k in ("zone", "sgg_name", "sgg", "umd", "biz", "stage", "stage_no", "households")})
    if a.dry_run or missing:
        return
    print(f"저장 {save(zones)}건")


if __name__ == "__main__":
    main()
